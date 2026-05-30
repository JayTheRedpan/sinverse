#!/usr/bin/env python3
"""
sinverse_story_import.py
========================

Local helper for bulk-importing Sinverse library stories from PDF and DOCX files.

WHAT IT DOES
------------
For every .pdf and .docx in the same folder as this script, it will:
  1. Extract the text.
  2. Repair PDF copy-paste damage:
       - join lines broken mid-sentence
       - keep real paragraph breaks
       - normalise smart quotes / em-dashes / odd spacing
  3. Write a clean Markdown file ("<name>.md") next to the PDF.
  4. Count the words.
  5. Append a starter entry to "library_entries.json" with the word count
     filled in and placeholders for everything you need to edit by hand.

Your PDFs never leave your computer. This runs entirely locally.

REQUIREMENTS
------------
  Python 3.8+
  For PDFs - one of these (tried in order, uses whichever is installed):
       pdfplumber   (best quality)   ->  pip install pdfplumber
       pypdf                          ->  pip install pypdf
       PyPDF2                         ->  pip install PyPDF2

  For DOCX files:
       python-docx                    ->  pip install python-docx

  Install both recommended libraries with:
       pip install pdfplumber python-docx

HOW TO USE
----------
  1. Put this script in a folder.
  2. Put the PDF and/or DOCX files you want to import in the SAME folder.
  3. Run:   python sinverse_story_import.py
  4. Collect the generated .md files and library_entries.json.
  5. Open library_entries.json, fill in the placeholder fields
     (title, author, tags, synopsis, dates, characters, canonical),
     then paste the entries into your real library.json.
  6. Move the .md files into your repo's library/stories/ folder.

The script never overwrites an existing .md file - if "story.md" already
exists it skips that PDF so you don't lose edits.
"""

import os
import re
import sys
import json

# --------------------------------------------------------------------------
# PDF text extraction - try whichever library is available
# --------------------------------------------------------------------------
def extract_text(pdf_path):
    # Try pdfplumber first (best layout handling)
    try:
        import pdfplumber
        text_parts = []
        with pdfplumber.open(pdf_path) as pdf:
            for page in pdf.pages:
                t = page.extract_text() or ""
                text_parts.append(t)
        return "\n\n".join(text_parts)
    except ImportError:
        pass

    # Fall back to pypdf
    try:
        from pypdf import PdfReader
        reader = PdfReader(pdf_path)
        return "\n\n".join((p.extract_text() or "") for p in reader.pages)
    except ImportError:
        pass

    # Fall back to PyPDF2
    try:
        from PyPDF2 import PdfReader
        reader = PdfReader(pdf_path)
        return "\n\n".join((p.extract_text() or "") for p in reader.pages)
    except ImportError:
        pass

    print("\nERROR: No PDF library found. Install one with:")
    print("    pip install pdfplumber\n")
    sys.exit(1)


def extract_docx(docx_path):
    try:
        import docx  # python-docx
    except ImportError:
        print("\nERROR: python-docx not found. Install it with:")
        print("    pip install python-docx\n")
        sys.exit(1)
    document = docx.Document(docx_path)
    # Each paragraph in the docx becomes its own line; blank paragraphs
    # act as paragraph separators. clean_text() handles the rest.
    parts = []
    for para in document.paragraphs:
        parts.append(para.text)
    return "\n".join(parts)


# --------------------------------------------------------------------------
# Text cleanup - repair PDF copy-paste damage
# --------------------------------------------------------------------------
def clean_text(raw):
    if not raw:
        return ""

    text = raw

    # Normalise line endings
    text = text.replace("\r\n", "\n").replace("\r", "\n")

    # Fix common PDF character artifacts
    replacements = {
        "\u201c": '"', "\u201d": '"',   # smart double quotes
        "\u2018": "'", "\u2019": "'",   # smart single quotes
        "\u2013": "-", "\u2014": "-",   # en/em dashes -> hyphen (optional)
        "\u2026": "...",                 # ellipsis
        "\u00a0": " ",                   # non-breaking space
        "\ufb01": "fi", "\ufb02": "fl",  # ligatures
        "\t": " ",
    }
    for bad, good in replacements.items():
        text = text.replace(bad, good)

    # De-hyphenate words split across line breaks: "transfor-\nmation" -> "transformation"
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)

    # Collapse 3+ blank lines down to a paragraph break (2 newlines)
    text = re.sub(r"\n{3,}", "\n\n", text)

    # Split into paragraphs on blank lines
    paragraphs = re.split(r"\n\s*\n", text)

    cleaned_paragraphs = []
    for para in paragraphs:
        # Within a paragraph, join the positional line breaks into one flowing line.
        lines = [ln.strip() for ln in para.split("\n") if ln.strip()]
        if not lines:
            continue
        joined = ""
        for i, line in enumerate(lines):
            if i == 0:
                joined = line
            else:
                # If previous text ends with sentence-ending punctuation, keep as
                # part of same paragraph but still join with a space (it's one para).
                joined = joined.rstrip() + " " + line
        # Collapse any double spaces produced by joining
        joined = re.sub(r" {2,}", " ", joined).strip()
        cleaned_paragraphs.append(joined)

    # Re-assemble with blank lines between paragraphs (proper Markdown)
    return "\n\n".join(cleaned_paragraphs).strip() + "\n"


# --------------------------------------------------------------------------
# Word count
# --------------------------------------------------------------------------
def count_words(text):
    return len(re.findall(r"\b\w+\b", text))


# --------------------------------------------------------------------------
# Build a slug for the filename / id
# --------------------------------------------------------------------------
def slugify(name):
    s = name.lower()
    s = re.sub(r"[^\w\s-]", "", s)
    s = re.sub(r"[\s_-]+", "_", s)
    return s.strip("_")


# --------------------------------------------------------------------------
# Main
# --------------------------------------------------------------------------
def main():
    here = os.path.dirname(os.path.abspath(__file__))
    sources = sorted(
        f for f in os.listdir(here)
        if f.lower().endswith(".pdf") or f.lower().endswith(".docx")
    )
    # Skip temporary Word lock files like ~$story.docx
    sources = [f for f in sources if not f.startswith("~$")]

    if not sources:
        print("No PDF or DOCX files found in this folder.")
        print("Put your .pdf / .docx files next to this script and run it again.")
        return

    print("Found %d file(s).\n" % len(sources))

    entries = []

    for src_name in sources:
        base = os.path.splitext(src_name)[0]
        ext  = os.path.splitext(src_name)[1].lower()
        slug = slugify(base)
        md_name = slug + ".md"
        md_path = os.path.join(here, md_name)
        src_path = os.path.join(here, src_name)

        print("Processing: %s" % src_name)

        # Don't clobber existing markdown
        if os.path.exists(md_path):
            print("  -> %s already exists, skipping extraction (keeping your edits)." % md_name)
            with open(md_path, "r", encoding="utf-8") as f:
                cleaned = f.read()
        else:
            if ext == ".pdf":
                raw = extract_text(src_path)
            elif ext == ".docx":
                raw = extract_docx(src_path)
            else:
                print("  -> unsupported file type, skipping.")
                continue
            cleaned = clean_text(raw)
            with open(md_path, "w", encoding="utf-8") as f:
                f.write(cleaned)
            print("  -> wrote %s" % md_name)

        wc = count_words(cleaned)
        print("  -> %d words\n" % wc)

        # Starter library.json entry with placeholders.
        # id is null - you'll assign real ids when merging into library.json.
        entry = {
            "id": None,                    # assign a real unique id when merging
            "type": "standalone",
            "title": base,                 # PLACEHOLDER - edit to a nice display title
            "author": "REPLACE_AUTHOR",    # PLACEHOLDER - must match contributors.json id
            "summary": None,               # PLACEHOLDER - short blurb shown on the card
            "tags": [],                    # PLACEHOLDER - add content tags
            "file": "stories/" + md_name,
            "date": "REPLACE_YYYY-MM",     # PLACEHOLDER - posted date e.g. "2025-05"
            "universe_date": None,         # PLACEHOLDER - integer year vs TDay, or null
            "canonical": False,            # PLACEHOLDER - true/false
            "characters": [],              # PLACEHOLDER - lowercase character names
            "wordCount": wc                # auto-filled
        }
        entries.append(entry)

    out_path = os.path.join(here, "library_entries.json")
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(entries, f, indent=2, ensure_ascii=False)

    print("=" * 60)
    print("Done. Wrote %d markdown file(s) and library_entries.json" % len(entries))
    print()
    print("NEXT STEPS:")
    print("  1. Open library_entries.json and fill in the REPLACE_ fields")
    print("     and the null 'summary' fields.")
    print("  2. Set real, unique 'id' values (currently null) that don't")
    print("     collide with your existing library.json entries.")
    print("  3. Paste the entries into your library.json.")
    print("  4. Move the .md files into library/stories/ in your repo.")
    print()
    print("  Serials: the script makes one standalone entry per PDF. If a PDF")
    print("  is actually one chapter of a serial, group those entries by hand")
    print("  into a single 'type':'serial' object with a 'chapters' array.")
    print("=" * 60)


if __name__ == "__main__":
    main()
